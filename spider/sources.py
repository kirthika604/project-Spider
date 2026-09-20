"""User-provided sources: files, folders and links (section 16).

A file's evidence is its sheet and cell or its PDF page number, so a value
from a spreadsheet can be traced exactly like a value from a web page.
"""

from __future__ import annotations

import csv
import json
from dataclasses import dataclass, field
from pathlib import Path

from .crawl.frontier import domain_of
from .extract.parser import ParsedPage
from .standardize.names import normalise
from .store.db import jdump, now

TEXT_SUFFIXES = {".txt", ".md", ".rst", ".text"}
SUPPORTED = {".pdf", ".csv", ".tsv", ".xlsx", ".xls", ".json"} | TEXT_SUFFIXES


@dataclass
class SourceTest:
    source_id: str
    kind: str
    readable: bool
    fields_found: list[str] = field(default_factory=list)
    rows: int = 0
    note: str = ""


def register(conn, item) -> None:
    conn.execute(
        "INSERT OR REPLACE INTO sources(id,type,location,tier,authoritative_for,config,"
        "added_at) VALUES(?,?,?,?,?,?,?)",
        (item.id, item.type, item.location, item.tier, jdump(item.authoritative_for),
         jdump({"map": item.map, "ai_allowed": item.ai_allowed, "pages": item.pages,
                "language": item.language, "license": item.license}), now()))
    conn.commit()


def read_source(conn, spec, item, root: Path) -> SourceTest:
    """Read one non-web source into pages/fields so `build` can use it."""
    register(conn, item)
    if item.type == "api":
        return _read_api(conn, spec, item)
    if item.type == "feed":
        return _read_feed(conn, spec, item)
    path = Path(item.location)
    if not path.is_absolute():
        path = (root / item.location)
    if item.type == "folder" or path.is_dir():
        return _read_folder(conn, spec, item, path)
    if not path.exists():
        return SourceTest(item.id, item.type, False, note=f"{path} not found")
    reader = {"pdf": _read_pdf, "csv": _read_csv, "xlsx": _read_xlsx,
              "json": _read_json, "text": _read_text, "file": _read_text}
    suffix = path.suffix.lower()
    kind = item.type if item.type in reader else (
        "csv" if suffix in (".csv", ".tsv") else
        "xlsx" if suffix in (".xlsx", ".xls") else
        "pdf" if suffix == ".pdf" else
        "json" if suffix == ".json" else "text")
    return reader[kind](conn, spec, item, path)


def _crawl_row(conn, item) -> int:
    cursor = conn.execute(
        "INSERT INTO crawls(started_at,seeds,keywords,depth,max_pages,kind) "
        "VALUES(?,?,?,?,?,?)", (now(), jdump([item.location]), "[]", 0, 0, "source"))
    conn.commit()
    return cursor.lastrowid


def _save(conn, item, crawl_id, url, title, text, fields, tier=None) -> int:
    """Store a file the same way a page is stored, so evidence works alike."""
    page = ParsedPage(url=url, title=title, text=text)
    row = conn.execute("SELECT id FROM pages WHERE url=?", (url,)).fetchone()
    values = (crawl_id, url, domain_of(url) or "file", title, "", "", "", item.language or "",
              "", text, page.word_count, 1.0, 200, page.content_hash, 0,
              item.tier if tier is None else tier, item.id, now())
    if row:
        page_id = row["id"]
        conn.execute(
            "UPDATE pages SET crawl_id=?,url=?,domain=?,title=?,description=?,author=?,"
            "published=?,lang=?,headings=?,text=?,word_count=?,relevance=?,status=?,"
            "content_hash=?,depth=?,tier=?,source_id=?,fetched_at=? WHERE id=?",
            (*values, page_id))
        conn.execute("DELETE FROM fields WHERE page_id=?", (page_id,))
        conn.execute("DELETE FROM pages_fts WHERE rowid=?", (page_id,))
    else:
        page_id = conn.execute(
            "INSERT INTO pages(crawl_id,url,domain,title,description,author,published,"
            "lang,headings,text,word_count,relevance,status,content_hash,depth,tier,"
            "source_id,fetched_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
            values).lastrowid
    conn.executemany("INSERT INTO fields(page_id,name,value) VALUES(?,?,?)",
                     [(page_id, n, v) for n, v in fields])
    conn.execute("INSERT INTO pages_fts(rowid,title,description,headings,text) "
                 "VALUES(?,?,?,?,?)", (page_id, title, "", "", text))
    conn.execute("UPDATE sources SET pages_read = pages_read + 1 WHERE id=?", (item.id,))
    conn.commit()
    return page_id


# --------------------------------------------------------------- CSV / XLSX
def _column_map(item, spec, header: list[str]) -> dict[str, str]:
    """Map spreadsheet columns to schema fields, by `map:` or by name."""
    mapping = {}
    lowered = {str(h).strip().lower(): h for h in header if h}
    for column, field_name in (item.map or {}).items():
        if str(column).strip().lower() in lowered:
            mapping[lowered[str(column).strip().lower()]] = field_name
    known = {}
    for ent in spec.entities.values():
        for field_name in ent.fields:
            known[field_name.lower()] = field_name
            known[field_name.replace("_", " ").lower()] = field_name
    for low, original in lowered.items():
        if original in mapping:
            continue
        if low in known:
            mapping[original] = known[low]
    return mapping


def _read_csv(conn, spec, item, path: Path) -> SourceTest:
    delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
    with open(path, newline="", encoding="utf-8-sig") as handle:
        rows = list(csv.DictReader(handle, delimiter=delimiter))
        header = list(rows[0].keys()) if rows else []
    mapping = _column_map(item, spec, header)
    crawl_id = _crawl_row(conn, item)
    saved = 0
    for index, row in enumerate(rows, start=2):        # row 1 is the header
        fields = [(field_name, normalise(row.get(column)))
                  for column, field_name in mapping.items()
                  if normalise(row.get(column))]
        if not fields:
            continue
        lines = [f"{column}: {row.get(column)}" for column in header
                 if normalise(row.get(column))]
        text = f"Row {index} of {path.name}. " + ". ".join(lines) + "."
        title = next((normalise(row.get(c)) for c, f in mapping.items()
                      if f in (list(spec.entities.values())[0].identity or [])), "") \
            or f"{path.name} row {index}"
        url = f"file://{path.resolve()}#row={index}"
        _save(conn, item, crawl_id, url, title, text, fields)
        saved += 1
    conn.execute("UPDATE sources SET values_given = values_given + ? WHERE id=?",
                 (saved * max(1, len(mapping)), item.id))
    conn.commit()
    return SourceTest(item.id, "csv", True, sorted(set(mapping.values())), saved,
                      f"{saved} rows read from {path.name}")


def _read_xlsx(conn, spec, item, path: Path) -> SourceTest:
    try:
        from openpyxl import load_workbook
    except ImportError:
        return SourceTest(item.id, "xlsx", False,
                          note="reading Excel needs openpyxl "
                               "(pip install 'project-spider[files]')")
    book = load_workbook(path, data_only=True, read_only=True)
    crawl_id = _crawl_row(conn, item)
    saved, found_fields = 0, set()
    for sheet in book.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            continue
        header = [str(c) if c is not None else "" for c in rows[0]]
        mapping = _column_map(item, spec, header)
        if not mapping:
            continue
        for index, values in enumerate(rows[1:], start=2):
            record = dict(zip(header, values))
            fields = [(field_name, normalise(record.get(column)))
                      for column, field_name in mapping.items()
                      if normalise(record.get(column))]
            if not fields:
                continue
            found_fields.update(f for f, _ in fields)
            cells = ", ".join(
                f"{column} ({sheet.title}!{_cell(header.index(column), index)}): "
                f"{record.get(column)}"
                for column in mapping if normalise(record.get(column)))
            text = f"{path.name}, sheet {sheet.title}, row {index}. {cells}."
            identity_fields = list(spec.entities.values())[0].identity or []
            title = next((normalise(record.get(c)) for c, f in mapping.items()
                          if f in identity_fields), "") or f"{sheet.title} row {index}"
            url = f"file://{path.resolve()}#{sheet.title}!{index}"
            _save(conn, item, crawl_id, url, title, text, fields)
            saved += 1
    book.close()
    conn.execute("UPDATE sources SET values_given = values_given + ? WHERE id=?",
                 (saved, item.id))
    conn.commit()
    return SourceTest(item.id, "xlsx", True, sorted(found_fields), saved,
                      f"{saved} rows read from {path.name}")


def _cell(column_index: int, row_index: int) -> str:
    letters, number = "", column_index + 1
    while number:
        number, remainder = divmod(number - 1, 26)
        letters = chr(65 + remainder) + letters
    return f"{letters}{row_index}"


# ---------------------------------------------------------------- PDF / text
def _read_pdf(conn, spec, item, path: Path) -> SourceTest:
    pages = _pdf_pages(path)
    if pages is None:
        return SourceTest(item.id, "pdf", False,
                          note="reading PDFs needs pypdf "
                               "(pip install 'project-spider[files]')")
    wanted = _page_range(item.pages, len(pages))
    crawl_id = _crawl_row(conn, item)
    saved = 0
    for number, text in pages:
        if wanted and number not in wanted:
            continue
        if not text.strip():
            continue
        url = f"file://{path.resolve()}#page={number}"
        title = f"{path.name} p.{number}"
        _save(conn, item, crawl_id, url, title, text, [])
        saved += 1
    return SourceTest(item.id, "pdf", True, [], saved,
                      f"{saved} pages read from {path.name}; evidence carries the page number")


def _pdf_pages(path: Path):
    try:
        from pypdf import PdfReader
    except ImportError:
        try:
            from PyPDF2 import PdfReader          # older name
        except ImportError:
            return None
    try:
        reader = PdfReader(str(path))
    except Exception:
        return []
    out = []
    for index, page in enumerate(reader.pages, start=1):
        try:
            out.append((index, page.extract_text() or ""))
        except Exception:
            out.append((index, ""))
    return out


def _page_range(text: str | None, total: int):
    if not text:
        return None
    wanted = set()
    for part in str(text).split(","):
        part = part.strip()
        if "-" in part:
            start, _, end = part.partition("-")
            try:
                wanted.update(range(int(start), min(int(end), total) + 1))
            except ValueError:
                continue
        elif part.isdigit():
            wanted.add(int(part))
    return wanted or None


def _read_text(conn, spec, item, path: Path) -> SourceTest:
    try:
        text = path.read_text(encoding="utf-8", errors="replace")
    except OSError as exc:
        return SourceTest(item.id, "text", False, note=str(exc))
    crawl_id = _crawl_row(conn, item)
    _save(conn, item, crawl_id, f"file://{path.resolve()}", path.stem, text, [])
    return SourceTest(item.id, "text", True, [], 1, f"read {path.name}")


def _read_json(conn, spec, item, path: Path) -> SourceTest:
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError) as exc:
        return SourceTest(item.id, "json", False, note=str(exc))
    records = data if isinstance(data, list) else data.get("records", [data])
    crawl_id = _crawl_row(conn, item)
    saved = 0
    for index, record in enumerate(records, start=1):
        if not isinstance(record, dict):
            continue
        mapping = _column_map(item, spec, list(record))
        fields = [(f, normalise(record.get(c))) for c, f in mapping.items()
                  if normalise(record.get(c))]
        text = f"{path.name} record {index}. " + ". ".join(
            f"{k}: {v}" for k, v in record.items() if v not in (None, "")) + "."
        _save(conn, item, crawl_id, f"file://{path.resolve()}#record={index}",
              f"{path.stem} {index}", text, fields)
        saved += 1
    return SourceTest(item.id, "json", True, [], saved, f"{saved} records read")


def _read_folder(conn, spec, item, path: Path) -> SourceTest:
    if not path.exists():
        return SourceTest(item.id, "folder", False, note=f"{path} not found")
    read, skipped, fields = 0, [], set()
    from .spec import SourceItem
    for child in sorted(path.rglob("*")):
        if child.is_dir():
            continue
        if child.suffix.lower() not in SUPPORTED:
            skipped.append(child.name)
            continue
        child_item = SourceItem(
            id=f"{item.id}:{child.name}", type="file", location=str(child),
            tier=item.tier, authoritative_for=item.authoritative_for,
            map=item.map, ai_allowed=item.ai_allowed, language=item.language)
        result = read_source(conn, spec, child_item, path)
        read += 1
        fields.update(result.fields_found)
    note = f"{read} files read"
    if skipped:
        note += f"; skipped {len(skipped)} unsupported: {', '.join(skipped[:5])}"
    return SourceTest(item.id, "folder", True, sorted(fields), read, note)


def _read_api(conn, spec, item) -> SourceTest:
    """A JSON endpoint the user knows, with its key (section 16).

    The key is read from the environment, never written into spider.yaml.
    """
    import os

    import requests

    from . import USER_AGENT
    config = item.map or {}
    headers = {"User-Agent": USER_AGENT, "Accept": "application/json"}
    key_name = getattr(item, "key_env", None) or config.pop("key_env", None)
    if key_name:
        key = os.environ.get(str(key_name))
        if not key:
            return SourceTest(item.id, "api", False,
                              note=f"{key_name} is not set - add it to .env")
        headers["Authorization"] = f"Bearer {key}"
    try:
        response = requests.get(item.location, headers=headers, timeout=20)
    except Exception as exc:
        return SourceTest(item.id, "api", False, note=f"{type(exc).__name__}: {exc}"[:120])
    if response.status_code != 200:
        return SourceTest(item.id, "api", False,
                          note=f"the endpoint answered HTTP {response.status_code}")
    try:
        data = response.json()
    except ValueError:
        return SourceTest(item.id, "api", False, note="the endpoint did not return JSON")

    records = data if isinstance(data, list) else (
        data.get("records") or data.get("results") or data.get("data") or [data])
    if isinstance(records, dict):
        records = [records]
    crawl_id = _crawl_row(conn, item)
    saved, found = 0, set()
    for index, record in enumerate(records, start=1):
        if not isinstance(record, dict):
            continue
        flat = _flatten(record)
        mapping = _column_map(item, spec, list(flat))
        fields = [(field_name, normalise(flat.get(column)))
                  for column, field_name in mapping.items()
                  if normalise(flat.get(column))]
        found.update(f for f, _ in fields)
        text = (f"{item.location}, record {index}. " +
                ". ".join(f"{k}: {v}" for k, v in flat.items()
                          if v not in (None, "", [])) + ".")
        title = next((normalise(flat.get(c)) for c, f in mapping.items()
                      if f in (list(spec.entities.values())[0].identity or [])), "") \
            or f"{item.id} record {index}"
        _save(conn, item, crawl_id, f"{item.location}#record={index}", title, text, fields)
        saved += 1
    return SourceTest(item.id, "api", True, sorted(found), saved,
                      f"{saved} record(s) read from the endpoint")


def _flatten(record: dict, prefix: str = "") -> dict:
    """Nested JSON becomes `parent.child` keys, so `map:` can reach them."""
    out = {}
    for key, value in record.items():
        name = f"{prefix}{key}"
        if isinstance(value, dict):
            out.update(_flatten(value, f"{name}."))
        elif isinstance(value, list):
            out[name] = "; ".join(str(v) for v in value if not isinstance(v, (dict, list)))
        else:
            out[name] = value
    return out


def _read_feed(conn, spec, item) -> SourceTest:
    """An RSS or Atom feed: its entries become pages to crawl."""
    import xml.etree.ElementTree as ElementTree

    import requests

    from . import USER_AGENT
    try:
        response = requests.get(item.location, timeout=20,
                                headers={"User-Agent": USER_AGENT})
        root_node = ElementTree.fromstring(response.content)
    except Exception as exc:
        return SourceTest(item.id, "feed", False,
                          note=f"cannot read the feed: {type(exc).__name__}: {exc}"[:120])

    links = []
    for entry in root_node.iter():
        tag = entry.tag.split("}")[-1]
        if tag == "item":
            link = entry.findtext("link")
            if link:
                links.append(link.strip())
        elif tag == "entry":
            for child in entry:
                if child.tag.split("}")[-1] == "link" and child.get("href"):
                    links.append(child.get("href"))
                    break
    links = [url for url in dict.fromkeys(links) if url.startswith("http")]
    if not links:
        return SourceTest(item.id, "feed", False, note="the feed has no entry links")

    from .crawl.crawler import Crawler
    rules: dict[str, list[str]] = {}
    for ent in spec.entities.values():
        for name, fld in ent.fields.items():
            if fld.extract:
                rules.setdefault(name, []).extend(fld.extract)
    crawler = Crawler(conn, keywords=spec.sources.keywords, depth=0,
                      max_pages=len(links), delay=spec.sources.delay_seconds,
                      any_domain=True, extract_rules=rules,
                      tier_of=lambda _url: item.tier, source_id=item.id)
    result = crawler.run(links, kind="source")
    return SourceTest(item.id, "feed", True, [], result.saved,
                      f"{len(links)} entry link(s) in the feed, {result.saved} saved")


def health(conn) -> list[dict]:
    return [dict(r) for r in conn.execute(
        "SELECT s.id, s.type, s.location, s.tier, s.pages_read, "
        "  (SELECT COUNT(*) FROM attributes a JOIN pages p ON p.id=a.source_page "
        "   WHERE p.source_id=s.id AND a.status='accepted') AS values_accepted, "
        "  (SELECT COUNT(*) FROM attributes a JOIN pages p ON p.id=a.source_page "
        "   WHERE p.source_id=s.id AND a.status='review') AS values_in_review "
        "FROM sources s ORDER BY values_accepted DESC")]
